import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from collections import OrderedDict

wb = openpyxl.Workbook()

# ═══ FILLS ═══
DARK_BLUE, MED_BLUE, LIGHT_BLUE = "1F3864", "2E75B6", "D6E4F0"
ACCENT_GREEN, ACCENT_RED, ACCENT_ORANGE = "548235", "C00000", "ED7D31"
BORDER_COLOR = "B4C6E7"
BG_ALT_ROW = "F2F7FB"

hdr_fill     = PatternFill("solid", fgColor=DARK_BLUE)
alt_fill     = PatternFill("solid", fgColor=BG_ALT_ROW)
pass_fill    = PatternFill("solid", fgColor="E2EFDA")
fail_fill    = PatternFill("solid", fgColor="FCE4EC")
warn_fill    = PatternFill("solid", fgColor="FFF2CC")
green_fill   = PatternFill("solid", fgColor="C6EFCE")
red_fill     = PatternFill("solid", fgColor="FFC7CE")
mod_fill_a   = PatternFill("solid", fgColor="EBF1F8")
white_fill   = PatternFill("solid", fgColor="FFFFFF")
orange_fill  = PatternFill("solid", fgColor=ACCENT_ORANGE)

hdr_font   = Font(bold=True, color="FFFFFF", size=11, name="Calibri")
title_font = Font(bold=True, color=DARK_BLUE, size=16, name="Calibri")
sub_font   = Font(color="666666", size=10, name="Calibri", italic=True)
norm_font  = Font(color="333333", size=10, name="Calibri")
bold_font  = Font(bold=True, color="333333", size=10, name="Calibri")
grn_font   = Font(bold=True, color=ACCENT_GREEN, size=10, name="Calibri")
red_font   = Font(bold=True, color=ACCENT_RED, size=10, name="Calibri")
note_font  = Font(color="888888", size=9, name="Calibri", italic=True)

thin_border = Border(
    left=Side('thin', BORDER_COLOR), right=Side('thin', BORDER_COLOR),
    top=Side('thin', BORDER_COLOR), bottom=Side('thin', BORDER_COLOR))
total_border = Border(
    top=Side('medium', DARK_BLUE), bottom=Side('medium', DARK_BLUE),
    left=Side('thin', BORDER_COLOR), right=Side('thin', BORDER_COLOR))

center = Alignment(horizontal="center", vertical="center")

# ═══ DATA ═══
data = []
def add(module, cls, methods):
    data.append((module, cls, methods))

add("Controller - Auth", "AuthenticationControllerTest", [
    "authenticate_success", "introspect_success", "refreshToken_success",
    "logout_success", "authenticate_userNotFound_fail",
    "authenticate_wrongPassword_fail", "introspect_invalidToken_fail",
    "refreshToken_expired_fail", "outboundAuthenticate_success"])
add("Controller - Auth", "PermissionControllerTest",
    ["create_success", "getAll_success", "delete_success"])
add("Controller - Auth", "RoleControllerTest",
    ["create_success", "getAll_success", "delete_success"])
add("Controller - Auth", "UserControllerTest", [
    "CreateUser_validRequest_success", "CreateUser_usernameInvalid_fail",
    "CreateUser_passwordInvalid_fail", "CreateUser_userExisted_fail",
    "GetUsers_validRequest_success", "GetMyInfo_validRequest_success",
    "GetUserById_validRequest_success", "GetUserById_userNotFound_fail",
    "UpdateUser_validRequest_success", "GetMyInfo_userNotFound_fail",
    "UpdateUser_userNotFound_fail", "UpdateUser_notOwner_fail", "DeleteUser_success"])
add("Controller - Auth", "UserControllerIntegrationTest",
    [("CreateUser_validRequest_success", "Error", "Requires Docker / TestContainers")])
add("Controller - Community", "AdminCommunityControllerTest", [
    "adminDeletePost_success", "adminDeleteComment_success",
    "adminDeletePost_forbidden_whenNotAdmin", "adminDeleteComment_forbidden_whenNotAdmin"])
add("Controller - Community", "CommunityControllerTest", [
    "createPost_success", "getFeed_success", "getPost_success",
    "getPost_notFound_fail", "getUserPosts_success", "deletePost_success",
    "deletePost_notOwner_fail", "createComment_success", "getComments_success",
    "deleteComment_success", "likePost_success", "unlikePost_success",
    "isLiked_true", "isLiked_false", "getFollowingFeed_success",
    "follow_success", "unfollow_success", "getFollowing_success",
    "getFollowers_success", "getFollowingCount_success", "getFollowerCount_success",
    "isFollowing_success", "savePost_success", "unsavePost_success",
    "getSavedPosts_success", "isSaved_true", "isSaved_false", "uploadImage_success"])
add("Controller - Learning", "CourseControllerTest", [
    "createStep_success", "createProblem_success", "createLesson_success",
    "createChapter_success", "createLanguage_success", "getStep_success",
    "getProblem_success", "getLesson_success", "getChapter_success",
    "getLanguage_success", "getStepsByLessonForAdmin_success",
    "getLanguages_success", "updateStep_success", "deleteStep_success", "purgeStep_success"])
add("Controller - Learning", "EnrollmentControllerTest",
    ["enrollCourse_success", "enrollCourse_languageNotFound_fail"])
add("Controller - Learning", "PistonAPIControllerTest", ["runTest_success"])
add("Controller - Learning", "UserLearningControllerTest", [
    "getStepByLessonId_success", "getProblem_success",
    "getLessonByChapterId_success", "getChapterByLanguageName_success",
    "getLearningProgress_success", "verifyLesson_success",
    "getMyLearning_success", "submitPractice_success",
    "completePractice_success", "resetProgress_success"])
add("Controller - Other", "ChatControllerTest",
    ["chat_success", "getChat_success", "getChatOfUser_success"])
add("Service - AI", "ChatServiceTest",
    ["getHistorys_success", "getHistoryByUser_success"])
add("Service - Auth", "AuthenticationServiceTest", [
    "authenticate_userNotFound_fail", "authenticate_passwordWrong_fail",
    "authenticate_validRequest_success", "introspect_validToken_success",
    "introspect_invalidToken_fail", "logout_success",
    "refresh_token_success", "outboundAuthenticate_success"])
add("Service - Auth", "PermissionServiceTest",
    ["create_success", "getAll_success", "delete_success"])
add("Service - Auth", "RoleServiceTest",
    ["create_success", "getAll_success", "delete_success"])
add("Service - Auth", "UserServiceTest", [
    "CreateUser_validRequest_success", "CreateUser_userExisted_fail",
    "CreateUser_emailExisted_fail", "RoleNotFound_fail",
    "CreateUser_duplicateRaceConditions_fail", "UpdateUser_validRequest_fail",
    "UpdateUser_validRequest_success", "UpdateUser_owner_success",
    "UpdateUser_notOwner_fail", "updatePassword_wrongPassword_fail",
    "updatePassword_success", "getMyInfo_valid_success",
    "getMyInfo_userNotFound_error", "getUser_validRequest_success",
    "getUserById_validRequest_success", "getUserById_userNotFound_fail",
    "deleteUserById_success", "updateAvatar_success"])
add("Service - Cloudinary", "CloudinaryServiceTest",
    ["uploadFile_success", "uploadFile_ioException_fail"])
add("Service - Community", "CommentServiceTest", [
    "createComment_success", "createComment_postNotFound_fail",
    "createComment_userNotFound_fail", "getCommentsByPost_success",
    "deleteComment_success", "deleteComment_notOwner_fail", "deleteComment_notFound_fail"])
add("Service - Community", "FollowServiceTest",
    ["follow_self_fail", "follow_alreadyExists_fail", "follow_success"])
add("Service - Community", "PostLikeServiceTest",
    ["likePost_success", "likePost_alreadyLiked_fail", "unlikePost_notLiked_fail"])
add("Service - Community", "PostServiceTest",
    ["createPost_success", "getPostById_notFound_fail", "deletePost_notOwner_fail"])
add("Service - Community", "SavedPostServiceTest", [
    "savePost_success", "savePost_alreadyExists_fail", "unsavePost_success",
    "unsavePost_notFound_fail", "getSavedPosts_success", "isSaved_true", "isSaved_false"])
add("Service - Learning", "CourseServiceTest", [
    "getStepById_success", "getStepById_notFound_fail", "getProblemById_success",
    "getProblemById_notFound_fail", "getLessonById_success", "getChapterById_success",
    "getLanguageById_success", "getStepByLesson_success", "getProblemByLesson_success",
    "getLessonByChapter_success", "getChapterByLanguage_success",
    "getMyLearning_success", "getLearningProgress_success"])
add("Service - Learning", "EnrollmentServiceTest", [
    "enrollCourse_success", "enrollCourse_languageNotFound_fail", "enrollCourse_userNotFound_fail"])
add("Service - Learning", "LearningProgressServiceTest", [
    "getCurrentUser_success", "getCurrentUser_unauthenticated_fail",
    "completeLesson_updatesXpAndStreak", "completeLesson_alreadyCompleted_skips",
    "initializeLearningProgressForLanguage_success", "restProgress_success"])
add("Service - Learning", "LearningServiceTest", [
    "verifyLesson_success", "verifyLesson_lessonNotFound_fail",
    "verifyLesson_stepMismatch_fail", "verifyLesson_duplicateStepId_fail",
    "verifyLesson_wrongAnswer_fail", "verifyLesson_lockedLesson_fail",
    "verifyLesson_alreadyCompleted_returnsZeroXp"])
add("Service - Learning", "PracticeServiceTest", [
    "summit_success", "summit_problemNotFound_fail", "complete_notAllPassed_fail"])
add("Service - System", "Judge0APIServiceTest",
    ["executePythonCode_error_whenNoServer"])


# ═══════════════════════════════════════════════════════════
# SHEET 1: TEST CASES
# ═══════════════════════════════════════════════════════════
ws = wb.active
ws.title = "Test Cases"
ws.sheet_properties.tabColor = DARK_BLUE

ws.merge_cells("A1:F1")
c = ws["A1"]
c.value = "Identity Service - Test Coverage Report"
c.font = Font(bold=True, color=DARK_BLUE, size=16, name="Calibri")
c.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 36

ws.merge_cells("A2:F2")
c = ws["A2"]
c.value = "Generated: May 24, 2026  |  29 Test Classes  |  184 Test Methods"
c.font = sub_font
c.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[2].height = 22

ws.row_dimensions[3].height = 6

headers = ["#", "Module", "Test Class", "Test Method", "Result", "Note"]
widths  = [5, 30, 34, 56, 12, 36]
for i, (h, w) in enumerate(zip(headers, widths), 1):
    c = ws.cell(row=4, column=i, value=h)
    c.font = hdr_font; c.fill = hdr_fill; c.alignment = center; c.border = thin_border
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[4].height = 24

row = 5; stt = 0; mod_ranges = {}
for module, cls, methods in data:
    for m in methods:
        stt += 1
        if isinstance(m, tuple):
            method_name, status, note = m
        else:
            method_name, status, note = m, "Pass", ""

        vals = [stt, module, cls, method_name, status, note]
        for col, val in enumerate(vals, 1):
            c = ws.cell(row=row, column=col, value=val)
            c.font = norm_font if col <= 4 else (grn_font if status == "Pass" else red_font)
            c.alignment = center if col in (1, 5) else Alignment(vertical="center")
            c.border = thin_border
            if col == 5:
                c.fill = pass_fill if status == "Pass" else fail_fill
            if col == 6 and note:
                c.font = note_font
        ws.row_dimensions[row].height = 20

        if module not in mod_ranges:
            mod_ranges[module] = [row, row]
        else:
            mod_ranges[module][1] = row
        row += 1

for idx, mod in enumerate(mod_ranges):
    fill = mod_fill_a if idx % 2 == 0 else white_fill
    for r in range(mod_ranges[mod][0], mod_ranges[mod][1] + 1):
        ws.cell(row=r, column=2).fill = fill

ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A4:F{row - 1}"
ws.page_setup.orientation = 'landscape'
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0


# ═══════════════════════════════════════════════════════════
# SHEET 2: EXECUTIVE SUMMARY
# ═══════════════════════════════════════════════════════════
ws2 = wb.create_sheet("Executive Summary")
ws2.sheet_properties.tabColor = ACCENT_GREEN

ws2.merge_cells("A1:D1")
c = ws2["A1"]
c.value = "Executive Summary"
c.font = Font(bold=True, color=DARK_BLUE, size=16, name="Calibri")
c.alignment = Alignment(horizontal="center", vertical="center")
ws2.row_dimensions[1].height = 36

ws2.merge_cells("A2:D2")
c = ws2["A2"]
c.value = "Test execution overview — Identity Service"
c.font = sub_font
c.alignment = Alignment(horizontal="center", vertical="center")
ws2.row_dimensions[2].height = 22

sum_headers = ["Module", "Methods", "Pass", "Error"]
sum_widths  = [32, 14, 14, 14]
for i, (h, w) in enumerate(zip(sum_headers, sum_widths), 1):
    c = ws2.cell(row=4, column=i, value=h)
    c.font = hdr_font; c.fill = hdr_fill; c.alignment = center; c.border = thin_border
    ws2.column_dimensions[get_column_letter(i)].width = w
ws2.row_dimensions[4].height = 24

modules = OrderedDict()
for module, cls, methods in data:
    modules.setdefault(module, {"methods": 0, "pass": 0, "error": 0})
    for m in methods:
        modules[module]["methods"] += 1
        if isinstance(m, tuple):
            modules[module]["error"] += 1
        else:
            modules[module]["pass"] += 1

r = 5; total_m = total_p = total_e = 0
for idx, (mod, s) in enumerate(modules.items()):
    total_m += s["methods"]; total_p += s["pass"]; total_e += s["error"]
    bg = alt_fill if idx % 2 == 0 else white_fill
    for col, val in enumerate([mod, s["methods"], s["pass"], s["error"]], 1):
        c = ws2.cell(row=r, column=col, value=val)
        c.font = norm_font if col <= 1 else bold_font
        c.alignment = center if col >= 2 else Alignment(vertical="center")
        c.border = thin_border; c.fill = bg
        if col == 3 and s["error"] == 0: c.fill = pass_fill
        if col == 4 and s["error"] > 0:
            c.fill = fail_fill; c.font = Font(bold=True, color=ACCENT_RED, size=10, name="Calibri")
    r += 1

for col, val in enumerate(["TOTAL", total_m, total_p, total_e], 1):
    c = ws2.cell(row=r, column=col, value=val)
    c.font = Font(bold=True, color=DARK_BLUE, size=11, name="Calibri")
    c.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    c.alignment = center if col >= 2 else Alignment(vertical="center")
    c.border = total_border
    if col == 3: c.fill = green_fill; c.font = Font(bold=True, color=ACCENT_GREEN, size=11, name="Calibri")
    if col == 4 and total_e > 0:
        c.fill = red_fill; c.font = Font(bold=True, color=ACCENT_RED, size=11, name="Calibri")
ws2.row_dimensions[r].height = 28; r += 2

# KPI Boxes
ws2.merge_cells(f"A{r}:B{r}")
ws2.cell(row=r, column=1, value="OVERALL PASS RATE").font = Font(bold=True, color="555555", size=10, name="Calibri")
ws2.cell(row=r, column=1).alignment = center; r += 1
pass_rate = round(total_p / total_m * 100, 1)
ws2.merge_cells(f"A{r}:B{r}")
c = ws2.cell(row=r, column=1, value=f"{pass_rate}%")
c.font = Font(bold=True, color=ACCENT_GREEN, size=28, name="Calibri")
c.alignment = center; c.fill = pass_fill; ws2.cell(row=r, column=2).fill = pass_fill
ws2.row_dimensions[r].height = 44

ws2.merge_cells(f"C{r-1}:D{r-1}")
ws2.cell(row=r-1, column=3, value="TEST EXECUTION").font = Font(bold=True, color="555555", size=10, name="Calibri")
ws2.cell(row=r-1, column=3).alignment = center
summaries = [("Total Tests", total_m, DARK_BLUE, LIGHT_BLUE),
             ("Passed", total_p, ACCENT_GREEN, "E2EFDA"),
             ("Failed / Error", total_e, ACCENT_RED if total_e > 0 else "888888",
              "FCE4EC" if total_e > 0 else "F2F2F2")]
for i, (label, val, clr, bg) in enumerate(summaries):
    r2 = r + i
    for col, v in enumerate([label, val], 3):
        c = ws2.cell(row=r2, column=col, value=v)
        c.font = Font(bold=True, color=clr, size=14 if col == 4 else 10, name="Calibri")
        c.alignment = center; c.border = thin_border
        c.fill = PatternFill("solid", fgColor=bg)
    ws2.row_dimensions[r2].height = 28

# Pie chart
r += 4
chart = PieChart(); chart.title = "Pass vs Error Distribution"
chart.style = 10; chart.width = 16; chart.height = 12
ws2.cell(row=r, column=1, value="Status"); ws2.cell(row=r, column=2, value="Count")
ws2.cell(row=r+1, column=1, value="Pass"); ws2.cell(row=r+1, column=2, value=total_p)
ws2.cell(row=r+2, column=1, value="Error"); ws2.cell(row=r+2, column=2, value=total_e)
chart.add_data(Reference(ws2, min_col=2, min_row=r, max_row=r+2), titles_from_data=True)
chart.set_categories(Reference(ws2, min_col=1, min_row=r+1, max_row=r+2))
s1 = DataPoint(idx=0); s1.graphicalProperties.solidFill = ACCENT_GREEN
s2 = DataPoint(idx=1); s2.graphicalProperties.solidFill = ACCENT_RED
chart.series[0].data_points = [s1, s2]
chart.dataLabels = DataLabelList()
chart.dataLabels.showPercent = True; chart.dataLabels.showVal = True
chart.dataLabels.showCatName = True
chart.dataLabels.font = Font(size=10, name="Calibri")
ws2.add_chart(chart, f"A{r+4}")
ws2.freeze_panes = "A4"


# ═══════════════════════════════════════════════════════════
# SHEET 3: COVERAGE DETAIL
# ═══════════════════════════════════════════════════════════
ws3 = wb.create_sheet("Coverage Detail")
ws3.sheet_properties.tabColor = MED_BLUE

ws3.merge_cells("A1:E1")
ws3["A1"].value = "Coverage Analysis by Layer"
ws3["A1"].font = Font(bold=True, color=DARK_BLUE, size=14, name="Calibri")
ws3["A1"].alignment = center
ws3.row_dimensions[1].height = 30

cov_headers = ["Layer", "Total", "Tested", "Missing", "Coverage"]
cov_data = [
    ("Controllers (@RestController)", 11, 11, 0, "100%"),
    ("Services (@Service)", 18, 17, 1, "94%"),
    ("Components (@Component)", 1, 0, 1, "0%"),
    ("ControllerAdvice", 1, 0, 1, "0%"),
    ("Repositories (@Repository)", 20, 0, 20, "0%"),
    ("Configurations (@Configuration)", 6, 0, 6, "0%"),
]
cov_widths = [34, 12, 12, 12, 14]
for i, (h, w) in enumerate(zip(cov_headers, cov_widths), 1):
    c = ws3.cell(row=3, column=i, value=h)
    c.font = hdr_font; c.fill = hdr_fill; c.alignment = center; c.border = thin_border
    ws3.column_dimensions[get_column_letter(i)].width = w

for i, (layer, total, tested, missing, pct) in enumerate(cov_data):
    r = 4 + i
    ws3.cell(row=r, column=1, value=layer).font = norm_font; ws3.cell(row=r, column=1).border = thin_border
    ws3.cell(row=r, column=2, value=total).font = norm_font; ws3.cell(row=r, column=2).alignment = center; ws3.cell(row=r, column=2).border = thin_border
    ws3.cell(row=r, column=3, value=tested).font = norm_font; ws3.cell(row=r, column=3).alignment = center; ws3.cell(row=r, column=3).border = thin_border
    ws3.cell(row=r, column=4, value=missing).font = norm_font; ws3.cell(row=r, column=4).alignment = center; ws3.cell(row=r, column=4).border = thin_border
    pc = ws3.cell(row=r, column=5, value=pct)
    pc.font = Font(bold=True, color=ACCENT_GREEN if pct == "100%" else (DARK_BLUE if "9" in pct else ACCENT_RED), size=11, name="Calibri")
    pc.alignment = center; pc.border = thin_border
    if pct == "100%": pc.fill = pass_fill
    elif "9" in pct: pc.fill = warn_fill
    else: pc.fill = fail_fill
    ws3.row_dimensions[r].height = 22

ws3.merge_cells("A11:E11")
ws3.cell(row=11, column=1, value="Items Not Yet Covered (Service + Controller scope)").font = Font(bold=True, color=DARK_BLUE, size=12, name="Calibri")
ws3.cell(row=11, column=1).alignment = center
ws3.row_dimensions[11].height = 28

miss_h = ["#", "Class", "Type", "Package", "Note"]
miss_w = [5, 38, 16, 30, 50]
for i, (h, w) in enumerate(zip(miss_h, miss_w), 1):
    c = ws3.cell(row=12, column=i, value=h)
    c.font = hdr_font; c.fill = orange_fill; c.alignment = center; c.border = thin_border
    ws3.column_dimensions[get_column_letter(i)].width = w

miss = [(1, "ChatServiceTest (main source)", "@Service", "service.ai", "Depends on Spring AI ChatClient; deferred"),
        (2, "UserSecurity", "@Component", "service.auth", "Out of scope (not service/controller)"),
        (3, "GlobalExceptionHandler", "@ControllerAdvice", "exception", "Out of scope (not service/controller)")]
for num, cls, typ, pkg, note in miss:
    r = 13 + num - 1
    ws3.cell(row=r, column=1, value=num).font = norm_font; ws3.cell(row=r, column=1).alignment = center; ws3.cell(row=r, column=1).border = thin_border
    ws3.cell(row=r, column=2, value=cls).font = norm_font; ws3.cell(row=r, column=2).border = thin_border
    ws3.cell(row=r, column=3, value=typ).font = norm_font; ws3.cell(row=r, column=3).alignment = center; ws3.cell(row=r, column=3).border = thin_border
    ws3.cell(row=r, column=4, value=pkg).font = norm_font; ws3.cell(row=r, column=4).border = thin_border
    ws3.cell(row=r, column=5, value=note).font = note_font; ws3.cell(row=r, column=5).border = thin_border

ws3.freeze_panes = "A4"

# ═══ SAVE ═══
out = "F:\\fisrt-backend\\my_project\\identity-servive\\Identity_Service_Test_Report.xlsx"
wb.save(out)
print(f"Saved: {out}")
print(f"Tests: {total_m} total | {total_p} Pass | {total_e} Error | Pass rate: {pass_rate}%")
